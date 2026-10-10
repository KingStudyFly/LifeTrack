#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将Markdown报告转换为docx格式"""

import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_run_font(run, font_name='微软雅黑', size=10.5, bold=False, color=None):
    """设置run字体"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

def parse_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # 设置默认中文字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        stripped = line.strip()

        # 跳过空行，但保留段落分隔
        if not stripped:
            i += 1
            continue

        # 标题
        if stripped.startswith('# ') and not stripped.startswith('## '):
            title = stripped[2:].strip()
            p = doc.add_heading(title, level=0)
            for run in p.runs:
                set_run_font(run, font_name='微软雅黑', size=18, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith('## '):
            title = stripped[3:].strip()
            p = doc.add_heading(title, level=1)
            for run in p.runs:
                set_run_font(run, font_name='微软雅黑', size=16, bold=True)
            i += 1
            continue

        if stripped.startswith('### '):
            title = stripped[4:].strip()
            p = doc.add_heading(title, level=2)
            for run in p.runs:
                set_run_font(run, font_name='微软雅黑', size=14, bold=True)
            i += 1
            continue

        if stripped.startswith('#### '):
            title = stripped[5:].strip()
            p = doc.add_heading(title, level=3)
            for run in p.runs:
                set_run_font(run, font_name='微软雅黑', size=12, bold=True)
            i += 1
            continue

        # 引用块（> 开头）
        if stripped.startswith('>'):
            text = stripped[1:].strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            set_run_font(run, font_name='微软雅黑', size=10, color=RGBColor(0x66, 0x66, 0x66))
            i += 1
            continue

        # 表格
        if stripped.startswith('|') and stripped.endswith('|'):
            # 收集表格所有行
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            # 过滤分隔行（如 |---|---|）
            data_lines = [l for l in table_lines if not re.match(r'^\|[\s\-:]+\|', l)]
            if len(data_lines) < 1:
                continue

            # 解析列
            def parse_row(row):
                parts = row[1:-1].split('|')
                return [p.strip() for p in parts]

            rows_data = [parse_row(l) for l in data_lines]
            if not rows_data:
                continue

            num_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=len(rows_data), cols=num_cols)
            table.style = 'Table Grid'

            for r_idx, row_data in enumerate(rows_data):
                row = table.rows[r_idx]
                for c_idx, cell_text in enumerate(row_data):
                    if c_idx < num_cols:
                        cell = row.cells[c_idx]
                        cell.text = cell_text
                        # 设置字体
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_run_font(run, font_name='微软雅黑', size=9)
            continue

        # 代码块
        if stripped.startswith('```'):
            lang = stripped[3:].strip()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1  # skip closing ```
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            run = p.add_run(code_text)
            set_run_font(run, font_name='Consolas', size=9, color=RGBColor(0x33, 0x33, 0x33))
            continue

        # 列表项
        list_match = re.match(r'^(\s*)[-*+]\s+(.*)', line)
        if list_match:
            indent_level = len(list_match.group(1)) // 2
            text = list_match.group(2).strip()
            p = doc.add_paragraph(style='List Bullet')
            # 清理markdown内联格式
            text = clean_inline_md(text)
            run = p.add_run(text)
            set_run_font(run, font_name='微软雅黑', size=10.5)
            i += 1
            continue

        # 普通段落
        text = clean_inline_md(stripped)
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_run_font(run, font_name='微软雅黑', size=10.5)
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")

def clean_inline_md(text):
    """清理简单的markdown内联格式"""
    # 加粗 **text**
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    # 斜体 *text*
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    # 行内代码 `text`
    text = re.sub(r'`(.+?)`', r'\1', text)
    # 链接 [text](url)
    text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
    return text

if __name__ == '__main__':
    md_file = r'D:\Coding\AISummary\LifeTrack\情绪可视化调研报告_raw.md'
    docx_file = r'D:\Coding\AISummary\LifeTrack\日记软件心情趋势与热力图实现方式调研报告.docx'
    parse_md_to_docx(md_file, docx_file)
